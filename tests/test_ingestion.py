"""Document-level tests; no invoice verdicts or audit rules are involved."""
from __future__ import annotations

import io
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from dataclasses import replace
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pymupdf
from jsonschema import ValidationError

from grader.extract import Mailbox, pages_of
from harness.ingestion import ingest_batch
from harness.models import Document, Extraction, OCRRegion, Page
from harness.readers import DocumentReader, message_sources
from harness.store import SourceConflict, StateStore

ROOT = Path(__file__).resolve().parents[1]


def pdf_bytes(*texts: str, scanned: bool = False) -> bytes:
    with pymupdf.open() as pdf:
        for text in texts:
            page = pdf.new_page(width=600, height=400)
            page.insert_text((36, 55), text, fontsize=18)
        if not scanned:
            return pdf.tobytes()
        with pymupdf.open() as images:
            for page in pdf:
                image = page.get_pixmap(dpi=160).tobytes("png")
                images.new_page(width=600, height=400).insert_image(page.rect, stream=image)
            return images.tobytes()


def xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Unfamiliar carrier"
    sheet.append(["Parcel", "Weight", "Amount"])
    sheet.append(["RIVER-908", 12.5, 123.45])
    sheet["B2"].number_format = "0.0"
    sheet["C2"].number_format = '"USD "0.00'
    hidden = workbook.create_sheet("Details")
    hidden.sheet_state = "hidden"
    hidden.append(["Hidden sheets are evidence too"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def email_bytes(
    *attachments: tuple[str, bytes],
    body: str = "Please see the enclosed invoice. Keep every reference exactly as received.",
    date: str = "Mon, 01 Jun 2026 10:00:00 +0000",
    html: bool = False,
) -> bytes:
    message = EmailMessage()
    message["From"] = "New Carrier <billing@new-carrier.example>"
    message["To"] = "Customer <ap@unseen.example>"
    message["Cc"] = "Operations <ops@unseen.example>"
    message["Date"] = date
    message["Subject"] = "Unseen layout and carrier"
    message["Message-ID"] = "<deliberately-reused-id@new-carrier.example>"
    message["X-GCA-Canary"] = "must not enter extracted text"
    message.set_content(body, subtype="html" if html else "plain")
    for filename, data in attachments:
        subtype = {
            ".pdf": "pdf",
            ".xlsx": "vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }.get(Path(filename).suffix.lower(), "octet-stream")
        message.add_attachment(data, maintype="application", subtype=subtype, filename=filename)
    return message.as_bytes()


class FakeOCR:
    cache_key = "test-ocr:1"

    def __init__(self) -> None:
        self.calls = 0
        # Intentionally out of order, including slightly offset cells on one row.
        self.regions = [
            OCRRegion("123.45", 0.95, [[150, 51], [215, 51], [215, 71], [150, 71]]),
            OCRRegion("RIVER-908", 0.98, [[20, 10], [140, 10], [140, 30], [20, 30]]),
            OCRRegion("Total", 0.97, [[20, 50], [90, 50], [90, 70], [20, 70]]),
        ]

    def recognize(self, image: bytes) -> list[OCRRegion]:
        assert image.startswith(b"\x89PNG")
        self.calls += 1
        return self.regions


class IngestionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.batch = self.root / "mailbox" / "batch_01"
        self.batch.mkdir(parents=True)
        self.state = self.root / "state"
        self.ocr = FakeOCR()
        self.reader = DocumentReader(self.ocr)

    def add_email(self, name: str = "0001.eml", **kwargs) -> bytes:
        data = email_bytes(**kwargs)
        (self.batch / name).write_bytes(data)
        return data

    def ingest(self):
        return ingest_batch(self.batch, self.state, reader=self.reader)

    def snapshot(self) -> str:
        with StateStore(self.state) as store:
            return "\n".join(store.connection.iterdump())

    def test_native_readers_match_grader_and_keep_exact_evidence(self) -> None:
        payload = email_bytes(("Frais été 001.PDF", pdf_bytes("Invoice RIVER-908", "Total USD 123.45")),
                              ("Tariff v9.xlsx", xlsx_bytes()))
        (self.batch / "0001.eml").write_bytes(payload)
        summary = self.ingest()
        self.assertEqual(summary.documents_added, 3)
        self.assertEqual(summary.problems, [])
        mailbox = Mailbox(self.batch.parent)
        with StateStore(self.state) as store:
            for document in store.documents():
                source = document.source
                raw, ext, canonical = mailbox.resolve(source.doc_ref)
                self.assertEqual(source.doc_ref, canonical)
                self.assertEqual(store.raw_bytes(source.doc_ref), raw)
                self.assertEqual([p.text for p in document.extraction.pages], pages_of(raw, ext))
                self.assertTrue(all(p.text == p.native_text for p in document.extraction.pages))
                for page in document.extraction.pages:
                    self.assertIsNone(mailbox.check(source.doc_ref, page.number, page.native_text))
                document.validate()
            email = store.get_document("batch_01/0001.eml")
            self.assertNotIn("must not enter extracted text", email.extraction.pages[0].text)
            self.assertEqual(email.source.metadata["from"][0]["address"], "billing@new-carrier.example")
            workbook = store.get_document("batch_01/0001.eml#Tariff v9.xlsx")
            self.assertEqual(len(workbook.extraction.pages), 2)
            self.assertIn("USD 123.45", workbook.extraction.pages[0].text)
            self.assertIn("Hidden sheets", workbook.extraction.pages[1].text)
        self.assertEqual(self.ocr.calls, 0)

    def test_repeated_batch_does_not_change_records_or_reextract(self) -> None:
        (self.batch / "0001.eml").write_bytes(email_bytes(("invoice.pdf", pdf_bytes("Invoice RIVER-908"))))
        first = self.ingest()
        before = self.snapshot()
        with patch.object(self.reader, "extract", side_effect=AssertionError("must use cache")) as extract:
            second = self.ingest()
        extract.assert_not_called()
        self.assertEqual(first.documents_added, 2)
        self.assertEqual(second.documents_added, 0)
        self.assertEqual(second.documents_existing, 2)
        self.assertEqual(second.extractions_reused, 2)
        self.assertEqual(self.snapshot(), before)

    def test_resends_keep_distinct_sources_but_share_content(self) -> None:
        data = email_bytes(("invoice.pdf", pdf_bytes("Invoice RIVER-908")))
        (self.batch / "0001.eml").write_bytes(data)
        (self.batch / "0002.eml").write_bytes(data)
        summary = self.ingest()
        self.assertEqual(summary.documents_added, 4)
        self.assertEqual(summary.extractions_run, 2)
        with StateStore(self.state) as store:
            self.assertEqual(store.counts(), {"batches": 1, "documents": 4, "blobs": 2, "extractions": 2, "pages": 2})
            first = store.get_source("batch_01/0001.eml#invoice.pdf")
            second = store.get_source("batch_01/0002.eml#invoice.pdf")
            self.assertNotEqual(first.parent_ref, second.parent_ref)
            self.assertEqual(first.content_hash, second.content_hash)

    def test_same_filename_different_bytes_remain_separate(self) -> None:
        for index in (1, 2):
            (self.batch / f"000{index}.eml").write_bytes(
                email_bytes(("invoice.pdf", pdf_bytes(f"Invoice NEW-{index}"))))
        self.ingest()
        with StateStore(self.state) as store:
            self.assertEqual(store.counts()["blobs"], 4)
            self.assertIn("NEW-1", store.get_document("batch_01/0001.eml#invoice.pdf").extraction.pages[0].text)
            self.assertIn("NEW-2", store.get_document("batch_01/0002.eml#invoice.pdf").extraction.pages[0].text)

    def test_state_survives_missing_input_and_future_batch(self) -> None:
        raw = self.add_email()
        self.ingest()
        self.batch.rename(self.batch.with_name("archived"))
        future = self.batch.with_name("batch_99")
        future.mkdir()
        (future / "0001.eml").write_bytes(email_bytes(date="Tue, 01 Dec 2026 10:00:00 +0000"))
        ingest_batch(future, self.state, reader=self.reader)
        with StateStore(self.state) as store:
            self.assertEqual(store.raw_bytes("batch_01/0001.eml"), raw)
            self.assertEqual(len(list(store.documents())), 2)
            self.assertEqual([d.source.doc_ref for d in store.documents(as_of="2026-06-01")], ["batch_01/0001.eml"])

    def test_received_time_normalized_without_changing_grader_text(self) -> None:
        self.add_email(date="Mon, 01 Jun 2026 00:30:00 +0530")
        self.ingest()
        with StateStore(self.state) as store:
            document = store.get_document("batch_01/0001.eml")
            self.assertEqual(document.source.received_at, "2026-05-31T19:00:00Z")
            self.assertIn("2026-06-01T00:30:00Z", document.extraction.pages[0].native_text)
            self.assertEqual(len(list(store.documents(as_of="2026-05-31"))), 1)
            self.assertEqual(len(list(store.documents(as_of="2026-05-30"))), 0)

    def test_changed_source_rolls_back_entire_batch(self) -> None:
        self.add_email("0002.eml")
        self.ingest()
        before = self.snapshot()
        self.add_email("0001.eml", body="New source processed before the conflict")
        self.add_email("0002.eml", body="Changed source at an existing reference")
        with self.assertRaisesRegex(SourceConflict, "source changed"):
            self.ingest()
        self.assertEqual(self.snapshot(), before)

    def test_ambiguous_attachment_names_fail_atomically(self) -> None:
        self.add_email("0001.eml")
        data = pdf_bytes("Invoice NEW-908")
        (self.batch / "0002.eml").write_bytes(email_bytes(("same.pdf", data), ("same.pdf", data)))
        with self.assertRaisesRegex(ValueError, "duplicate filename"):
            self.ingest()
        with StateStore(self.state) as store:
            self.assertTrue(all(count == 0 for count in store.counts().values()))

    def test_failed_extraction_is_persisted_and_retried(self) -> None:
        (self.batch / "0001.eml").write_bytes(email_bytes(("scan.pdf", pdf_bytes("Invoice RIVER-908", scanned=True))))
        with patch.object(self.ocr, "recognize", side_effect=RuntimeError("temporary OCR failure")):
            failed = self.ingest()
        self.assertEqual(len(failed.problems), 1)
        with StateStore(self.state) as store:
            doc = store.get_document("batch_01/0001.eml#scan.pdf")
            self.assertEqual(doc.extraction.status, "failed")
            self.assertEqual(doc.extraction.pages, [])
            self.assertTrue(store.raw_bytes(doc.source.doc_ref).startswith(b"%PDF"))
        retried = self.ingest()
        self.assertEqual(retried.problems, [])
        self.assertEqual(retried.documents_added, 0)
        self.assertEqual(retried.extractions_run, 1)
        with StateStore(self.state) as store:
            self.assertEqual(store.get_document("batch_01/0001.eml#scan.pdf").extraction.status, "ready")
            self.assertEqual(store.connection.execute("SELECT attempts FROM extractions WHERE format = 'pdf'").fetchone()[0], 2)
        before = self.snapshot()
        self.ingest()
        self.assertEqual(self.snapshot(), before)

    def test_corrupt_and_unsupported_files_are_visible_not_dropped(self) -> None:
        (self.batch / "0001.eml").write_bytes(email_bytes(("bad.pdf", b"broken pdf"), ("legacy.xyz", b"unknown")))
        summary = self.ingest()
        self.assertEqual(summary.documents_added, 3)
        self.assertEqual(len(summary.problems), 2)
        with StateStore(self.state) as store:
            self.assertEqual(store.get_document("batch_01/0001.eml#bad.pdf").extraction.status, "failed")
            self.assertEqual(store.get_document("batch_01/0001.eml#legacy.xyz").extraction.status, "unsupported")
            self.assertEqual(store.raw_bytes("batch_01/0001.eml#legacy.xyz"), b"unknown")

    def test_failed_shared_attachment_attempted_once_per_batch(self) -> None:
        raw = email_bytes(("scan.pdf", pdf_bytes("Invoice RIVER-908", scanned=True)))
        for name in ("0001.eml", "0002.eml"):
            (self.batch / name).write_bytes(raw)
        with patch.object(self.ocr, "recognize", side_effect=RuntimeError("OCR unavailable")) as ocr:
            summary = self.ingest()
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual(len(summary.problems), 2)
        self.assertEqual(summary.documents_added, 4)
        with StateStore(self.state) as store:
            self.assertEqual(store.connection.execute("SELECT attempts FROM extractions WHERE format = 'pdf'").fetchone()[0], 1)

    def test_attachment_filename_is_evidence_not_a_filesystem_path(self) -> None:
        filename = "../../outside.pdf"
        data = pdf_bytes("Invoice RIVER-908")
        (self.batch / "0001.eml").write_bytes(email_bytes((filename, data)))
        self.ingest()
        with StateStore(self.state) as store:
            self.assertEqual(store.raw_bytes(f"batch_01/0001.eml#{filename}"), data)
        self.assertEqual(list(self.root.rglob("outside.pdf")), [])

    def test_reader_version_invalidates_cache_without_duplicating_sources(self) -> None:
        self.add_email()
        self.ingest()
        old_version = self.reader.cache_key
        self.ocr.cache_key = "test-ocr:2"
        summary = self.ingest()
        self.assertEqual(summary.documents_added, 0)
        self.assertEqual(summary.extractions_run, 1)
        with StateStore(self.state) as store:
            doc = store.get_document("batch_01/0001.eml")
            self.assertEqual(doc.reader_version, self.reader.cache_key)
            self.assertIsNotNone(store.cached_extraction(doc.source.content_hash, "eml", old_version))
            self.assertEqual(store.counts()["documents"], 1)

    def test_html_only_email_keeps_native_text_and_readable_body(self) -> None:
        raw = self.add_email(body="<p>Invoice NEW-908</p><script>untrusted code</script><table><tr><td>Total</td><td>123.45</td></tr></table>", html=True)
        self.ingest()
        with StateStore(self.state) as store:
            page = store.get_document("batch_01/0001.eml").extraction.pages[0]
            self.assertEqual(page.native_text, pages_of(raw, "eml")[0])
            self.assertEqual(page.method, "html")
            self.assertIn("Invoice NEW-908", page.text)
            self.assertIn("Total 123.45", page.text)
            self.assertNotIn("untrusted code", page.text)

    def test_mixed_pdf_ocr_round_trip_preserves_native_pages(self) -> None:
        with pymupdf.open(stream=pdf_bytes("Invoice RIVER-908"), filetype="pdf") as mixed:
            with pymupdf.open(stream=pdf_bytes("Total 123.45", scanned=True), filetype="pdf") as scan:
                mixed.insert_pdf(scan)
            data = mixed.tobytes()
        (self.batch / "0001.eml").write_bytes(email_bytes(("mixed.pdf", data)))
        self.ingest()
        self.assertEqual(self.ocr.calls, 1)
        with StateStore(self.state) as store:
            doc = store.get_document("batch_01/0001.eml#mixed.pdf")
            native, scanned = doc.extraction.pages
            self.assertEqual(native.method, "grader")
            self.assertEqual(native.text, pages_of(data, "pdf")[0])
            self.assertEqual(scanned.number, 2)
            self.assertEqual(scanned.method, "ocr")
            self.assertEqual(scanned.native_text, "")
            self.assertEqual(scanned.text, "RIVER-908\nTotal 123.45")
            self.assertEqual(scanned.regions[0].confidence, 0.98)
            self.assertGreater(scanned.width, 0)
            self.assertGreater(scanned.height, 0)
            doc.validate()

    def test_scan_with_native_footer_still_runs_ocr(self) -> None:
        with pymupdf.open(stream=pdf_bytes("Invoice RIVER-908", scanned=True), filetype="pdf") as pdf:
            pdf[0].insert_text((36, 380), "Page 1")
            data = pdf.tobytes()
        result = self.reader.extract(data, "pdf")
        self.assertEqual(self.ocr.calls, 1)
        self.assertEqual(result.pages[0].native_text, "Page 1")
        self.assertEqual(result.pages[0].method, "ocr")

    def test_truly_blank_pdf_page_is_not_an_ocr_failure(self) -> None:
        with pymupdf.open() as pdf:
            pdf.new_page()
            data = pdf.tobytes()
        result = self.reader.extract(data, "pdf")
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.pages[0].text, "")
        self.assertEqual(self.ocr.calls, 0)

    def test_schema_rejects_invalid_provenance_and_page_numbers(self) -> None:
        self.add_email()
        source, _ = message_sources(self.batch, self.batch / "0001.eml")[0]
        valid = Document(source, "test:1", Extraction("ready", [Page(1, "text", "text")]))
        valid.validate()
        cases = [
            replace(valid, schema_version=2),
            replace(valid, source=replace(source, content_hash="invalid")),
            replace(valid, source=replace(source, doc_ref="another_batch/0001.eml")),
            replace(valid, source=replace(source, filename="invalid.pdf")),
            replace(valid, extraction=Extraction("ready", [Page(2, "text", "text")])),
            replace(valid, extraction=Extraction("failed")),
            replace(valid, extraction=Extraction("ready", [Page(1, "ocr", "", "ocr")])),
            replace(valid, extraction=Extraction("ready", [Page(
                1, "ocr", "", "ocr", [replace(self.ocr.regions[0], confidence=1.5)], 300, 400)])),
            replace(valid, extraction=Extraction("ready", [Page(
                1, "ocr", "", "ocr", [replace(self.ocr.regions[0], confidence=float("nan"))], 300, 400)])),
        ]
        for document in cases:
            with self.subTest(document=document), self.assertRaises((ValidationError, ValueError)):
                document.validate()

    def test_store_rejects_future_schema_and_untransactional_writes(self) -> None:
        self.add_email()
        source, data = message_sources(self.batch, self.batch / "0001.eml")[0]
        doc = Document(source, "test:1", Extraction("ready", [Page(1, "text", "text")]))
        with StateStore(self.state) as store:
            with self.assertRaisesRegex(RuntimeError, "transaction"):
                store.save(doc, data, extracted=True)
            with store.transaction(), self.assertRaisesRegex(ValueError, "hash"):
                store.save(doc, b"wrong bytes", extracted=True)
            self.assertEqual(store.connection.execute("PRAGMA foreign_keys").fetchone()[0], 1)
            store.connection.execute("PRAGMA user_version = 999")
        with self.assertRaisesRegex(ValueError, "unsupported state schema"):
            StateStore(self.state)

    def test_invalid_date_fails_without_partial_documents(self) -> None:
        raw = self.add_email()
        (self.batch / "0001.eml").write_bytes(raw.replace(b"Mon, 01 Jun 2026 10:00:00 +0000", b"not-a-date"))
        with self.assertRaisesRegex(ValueError, "Date header"):
            self.ingest()
        with StateStore(self.state) as store:
            self.assertEqual(store.counts()["documents"], 0)

    def test_cli_works_from_empty_cwd_across_fresh_processes(self) -> None:
        self.add_email()
        work = self.root / "empty"
        work.mkdir()
        command = [sys.executable, str(ROOT / "harness" / "harness.py"), "ingest", str(self.batch), "--state-dir", str(self.state), "--documents-only"]
        first = subprocess.run(command, cwd=work, capture_output=True, text=True, timeout=30)
        self.assertEqual(first.returncode, 0, first.stderr)
        before = self.snapshot()
        second = subprocess.run(command, cwd=work, capture_output=True, text=True, timeout=30)
        self.assertEqual(second.returncode, 0, second.stderr)
        self.assertEqual(second.stdout, "")
        self.assertIn("0 new, 1 existing", second.stderr)
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(list(work.iterdir()), [])
        with sqlite3.connect(self.state / "state.sqlite3") as db:
            self.assertEqual(db.execute("PRAGMA integrity_check").fetchone()[0], "ok")

    def test_cli_reports_problem_with_nonzero_exit(self) -> None:
        (self.batch / "0001.eml").write_bytes(email_bytes(("bad.pdf", b"bad")))
        result = subprocess.run(
            [sys.executable, str(ROOT / "harness" / "harness.py"), "ingest", str(self.batch), "--state-dir", str(self.state), "--documents-only"],
            cwd=self.root, capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("batch_01/0001.eml#bad.pdf", result.stderr)
        with StateStore(self.state) as store:
            self.assertEqual(store.counts()["documents"], 2)


class RealOCRTests(unittest.TestCase):
    def test_local_ocr_reads_unseen_scanned_invoice(self) -> None:
        data = pdf_bytes("INVOICE RIVER-908\nCustomer: UNSEEN LOGISTICS\nTotal USD 123.45", scanned=True)
        self.assertEqual(pages_of(data, "pdf"), [""])
        result = DocumentReader().extract(data, "pdf")
        self.assertEqual(result.status, "ready")
        page = result.pages[0]
        self.assertEqual(page.method, "ocr")
        self.assertIn("RIVER-908", page.text)
        self.assertIn("123.45", page.text)
        self.assertEqual(page.native_text, "")
        self.assertTrue(all(0 <= region.confidence <= 1 for region in page.regions))


if __name__ == "__main__":
    unittest.main()
